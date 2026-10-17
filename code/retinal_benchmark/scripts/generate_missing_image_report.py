#!/usr/bin/env python3
"""Generate a mirrored report of annotation rows whose images are missing."""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from retinal_benchmark.manifest import (  # noqa: E402
    DIAGNOSIS_LABELS,
    LESION_LABELS,
    build_image_index,
    choose_check_sheet,
    choose_lesion_sheet,
    codes_to_str,
    codes_to_text,
    extract_label_codes,
    normalize_key,
    worksheet_rows_by_patient,
)

IMAGE_ROOT = Path("/path/to/project/eyedataset/厦大AI")
INCOMING_ROOT = Path(
    "/path/to/project/eyedataset/_incoming_2026q1/2026-1月-3月平台标注"
)
OUTPUT_ROOT = PROJECT_ROOT / "reports" / "vision2_missing_images"


def main() -> None:
    image_index = build_image_index(IMAGE_ROOT)
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)

    workbook_summary = []
    unique_rows: dict[str, dict] = {}

    for workbook_path in sorted(INCOMING_ROOT.rglob("*.xlsx")):
        if workbook_path.name.startswith("~$") or workbook_path.name == ".xlsx":
            continue

        wb = load_workbook(workbook_path, read_only=True, data_only=True)
        sheet_names = list(wb.sheetnames)
        wb.close()

        lesion_sheet = choose_lesion_sheet(sheet_names)
        if lesion_sheet is None:
            continue
        check_sheet = choose_check_sheet(sheet_names, lesion_sheet)
        lesion_rows = worksheet_rows_by_patient(workbook_path, lesion_sheet)
        check_rows = worksheet_rows_by_patient(workbook_path, check_sheet) if check_sheet else {}

        missing_rows = []
        for patient_id, lesion_row in lesion_rows.items():
            image_key = normalize_key(patient_id)
            if image_key is None:
                continue
            matched_paths = sorted(image_index.get(image_key, []))
            if matched_paths:
                continue

            check_row = check_rows.get(patient_id, {})
            lesion_codes, diagnosis_codes, quality_codes = extract_label_codes(lesion_row, check_row)
            batch_id = workbook_path.parent.name.split("_", 1)[0]
            row = {
                "batch_id": batch_id,
                "patient_id": patient_id,
                "image_key": image_key,
                "source_annotation_file": str(workbook_path),
                "source_sheet": lesion_sheet,
                "quality_label": codes_to_str(quality_codes),
                "quality_text": "|".join(f"quality_{v}" for v in quality_codes),
                "diagnosis_ids": codes_to_str(diagnosis_codes),
                "diagnosis_texts": codes_to_text(diagnosis_codes, DIAGNOSIS_LABELS),
                "lesion_ids": codes_to_str(lesion_codes),
                "lesion_texts": codes_to_text(lesion_codes, LESION_LABELS),
            }
            missing_rows.append(row)
            unique_rows.setdefault(
                image_key,
                {
                    "batch_id": batch_id,
                    "patient_id": patient_id,
                    "image_key": image_key,
                    "first_annotation_file": str(workbook_path),
                    "annotation_files": set(),
                    "lesion_ids": set(),
                    "diagnosis_ids": set(),
                },
            )
            unique_rows[image_key]["annotation_files"].add(str(workbook_path))
            unique_rows[image_key]["lesion_ids"].update(int(v) for v in lesion_codes)
            unique_rows[image_key]["diagnosis_ids"].update(int(v) for v in diagnosis_codes)

        rel_dir = workbook_path.parent.relative_to(INCOMING_ROOT)
        out_dir = OUTPUT_ROOT / rel_dir
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / f"{workbook_path.stem}_missing_images.csv"
        with out_path.open("w", newline="", encoding="utf-8") as f:
            fieldnames = [
                "batch_id",
                "patient_id",
                "image_key",
                "source_annotation_file",
                "source_sheet",
                "quality_label",
                "quality_text",
                "diagnosis_ids",
                "diagnosis_texts",
                "lesion_ids",
                "lesion_texts",
            ]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(missing_rows)

        workbook_summary.append(
            {
                "batch_id": workbook_path.parent.name.split("_", 1)[0],
                "relative_dir": str(rel_dir),
                "workbook_name": workbook_path.name,
                "missing_rows": len(missing_rows),
                "report_csv": str(out_path),
            }
        )

    summary_path = OUTPUT_ROOT / "workbook_summary.csv"
    with summary_path.open("w", newline="", encoding="utf-8") as f:
        fieldnames = ["batch_id", "relative_dir", "workbook_name", "missing_rows", "report_csv"]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(workbook_summary)

    unique_path = OUTPUT_ROOT / "unique_missing_image_keys.csv"
    with unique_path.open("w", newline="", encoding="utf-8") as f:
        fieldnames = [
            "batch_id",
            "patient_id",
            "image_key",
            "first_annotation_file",
            "annotation_file_count",
            "annotation_files",
            "lesion_ids",
            "lesion_texts",
            "diagnosis_ids",
            "diagnosis_texts",
        ]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for image_key in sorted(unique_rows):
            item = unique_rows[image_key]
            lesion_ids = sorted(item["lesion_ids"])
            diagnosis_ids = sorted(item["diagnosis_ids"])
            writer.writerow(
                {
                    "batch_id": item["batch_id"],
                    "patient_id": item["patient_id"],
                    "image_key": item["image_key"],
                    "first_annotation_file": item["first_annotation_file"],
                    "annotation_file_count": len(item["annotation_files"]),
                    "annotation_files": json.dumps(sorted(item["annotation_files"]), ensure_ascii=False),
                    "lesion_ids": codes_to_str(lesion_ids),
                    "lesion_texts": codes_to_text(lesion_ids, LESION_LABELS),
                    "diagnosis_ids": codes_to_str(diagnosis_ids),
                    "diagnosis_texts": codes_to_text(diagnosis_ids, DIAGNOSIS_LABELS),
                }
            )

    stats = {
        "workbooks_scanned": len(workbook_summary),
        "workbooks_with_missing_rows": sum(1 for row in workbook_summary if row["missing_rows"] > 0),
        "missing_annotation_rows": sum(row["missing_rows"] for row in workbook_summary),
        "unique_missing_image_keys": len(unique_rows),
    }
    (OUTPUT_ROOT / "stats.json").write_text(
        json.dumps(stats, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    print(json.dumps(stats, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
