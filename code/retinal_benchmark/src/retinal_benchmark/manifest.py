"""Manifest building utilities."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from .constants import DATASET_SELECTIONS, DIAGNOSIS_LABELS, LESION_LABELS, VISION1_SELECTED_ANNOTATIONS


IMAGE_SUFFIX_PATTERN = re.compile(r"\.(png\.jpg|jpg|jpeg|png)$", re.IGNORECASE)


def normalize_key(text: str | None) -> str | None:
    if text is None:
        return None
    value = str(text).strip().replace(" ", "").replace("_", "")
    value = IMAGE_SUFFIX_PATTERN.sub("", value)
    return value


def parse_code_list(text: str | None) -> list[int]:
    if text is None:
        return []
    hits = re.findall(r"\d+", str(text))
    seen = []
    for hit in hits:
        value = int(hit)
        if value not in seen:
            seen.append(value)
    return seen


def safe_symlink(source: Path, target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists() or target.is_symlink():
        return
    os.symlink(str(source), str(target))


def assign_split(key: str) -> str:
    bucket = int(hashlib.md5(key.encode("utf-8")).hexdigest(), 16) % 100
    if bucket < 70:
        return "train"
    if bucket < 85:
        return "val"
    return "test"


def build_image_index(image_root: Path) -> dict[str, list[str]]:
    index: dict[str, list[str]] = defaultdict(list)
    for dirpath, _, filenames in os.walk(image_root):
        for filename in filenames:
            if not filename.lower().endswith((".jpg", ".jpeg", ".png")):
                continue
            key = normalize_key(filename)
            if key is None:
                continue
            index[key].append(str(Path(dirpath) / filename))
    return index


def worksheet_rows_by_patient(workbook_path: Path, sheet_name: str) -> dict[str, dict]:
    wb = load_workbook(workbook_path, read_only=False, data_only=True)
    if sheet_name not in wb.sheetnames:
        return {}
    ws = wb[sheet_name]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    header_to_col = {str(v): idx + 1 for idx, v in enumerate(headers) if v is not None}
    patient_col = header_to_col.get("Patient ID", 3)
    rows = {}
    for r in range(2, ws.max_row + 1):
        patient_id = ws.cell(r, patient_col).value
        if not patient_id:
            continue
        row = {}
        for header, col in header_to_col.items():
            row[header] = ws.cell(r, col).value
        rows[str(patient_id)] = row
    return rows


def choose_lesion_sheet(sheet_names: list[str]) -> str | None:
    if "序列维度" in sheet_names:
        return "序列维度"
    if "检查维度" in sheet_names:
        return "检查维度"
    return None


def choose_check_sheet(sheet_names: list[str], lesion_sheet: str | None) -> str | None:
    if "检查维度" in sheet_names:
        return "检查维度"
    return lesion_sheet


def extract_label_codes(lesion_row: dict, check_row: dict) -> tuple[list[int], list[int], list[int]]:
    lesion_codes = parse_code_list(
        lesion_row.get(
            "体征(1:屈光介质混浊;2:玻璃体出血;3:玻璃体星状小体;4:玻璃体后脱离;5:其他玻璃体异常（药物棒/气体/硅油）;6:玻璃膜疣;7:黄斑视网膜下纤维膜（黄斑盘变）;8:黄斑近视性萎缩斑;9:黄斑地图样萎缩;10:黄斑前膜;11:黄斑出血;12:黄斑视网膜下出血;13:黄斑区浆液性视网膜脱离;14:黄斑裂孔;15:其他黄斑病变（不在上述范围内）;16:视盘侧枝循环;17:视盘边界不清;18:先天性视盘发育异常;19:视盘大视杯（杯盘比≥03）;20:视神经萎缩;21:视盘新生血管;22:高度近视视盘萎缩弧;23:高度近视视盘萎缩环;24:视网膜脱离;25:视网膜出血象限性;26:视网膜下出血;27:视网膜前出血;28:视网膜裂孔;29:视网膜纤维膜;30:视网膜脉络膜占位（肿物）;31:视网膜激光斑（除外上述2种）;32:全视网膜光凝;33:单象限性视网膜光凝斑;34:视网膜骨细胞样色素改变;35:豹纹状眼底;36:晚霞状眼底;37:出血点、出血斑;38:硬性渗出;39:棉绒斑;40:视网膜新生血管;41:视网膜陈旧色素病灶;42:视网膜周边变性区;43:其他视网膜病变;44:视网膜动脉阻塞;45:视网膜血管白线;46:视网膜动脉硬化;47:视网膜血管鞘-视网膜血管炎)"
        )
        or check_row.get(
            "体征(1:屈光介质混浊;2:玻璃体出血;3:玻璃体星状小体;4:玻璃体后脱离;5:其他玻璃体异常（药物棒/气体/硅油）;6:玻璃膜疣;7:黄斑视网膜下纤维膜（黄斑盘变）;8:黄斑近视性萎缩斑;9:黄斑地图样萎缩;10:黄斑前膜;11:黄斑出血;12:黄斑视网膜下出血;13:黄斑区浆液性视网膜脱离;14:黄斑裂孔;15:其他黄斑病变（不在上述范围内）;16:视盘侧枝循环;17:视盘边界不清;18:先天性视盘发育异常;19:视盘大视杯（杯盘比≥03）;20:视神经萎缩;21:视盘新生血管;22:高度近视视盘萎缩弧;23:高度近视视盘萎缩环;24:视网膜脱离;25:视网膜出血象限性;26:视网膜下出血;27:视网膜前出血;28:视网膜裂孔;29:视网膜纤维膜;30:视网膜脉络膜占位（肿物）;31:视网膜激光斑（除外上述2种）;32:全视网膜光凝;33:单象限性视网膜光凝斑;34:视网膜骨细胞样色素改变;35:豹纹状眼底;36:晚霞状眼底;37:出血点、出血斑;38:硬性渗出;39:棉绒斑;40:视网膜新生血管;41:视网膜陈旧色素病灶;42:视网膜周边变性区;43:其他视网膜病变;44:视网膜动脉阻塞;45:视网膜血管白线;46:视网膜动脉硬化;47:视网膜血管鞘-视网膜血管炎)"
        )
    )
    diagnosis_codes = parse_code_list(
        check_row.get(
            "诊断(1:正常眼底;2:视网膜色素变性;3:视网膜中央静脉阻塞;4:视网膜分支静脉阻塞;5:结晶样视网膜营养不良（BCD）;6:FEVR (视网膜血管牵引);7:先天性脉络膜缺损;8:糖尿病视网膜病变（NPDR）;9:糖尿病视网膜病变（PDR）;10:湿性年龄相关黄斑变性;11:干性年龄相关黄斑变性;12:病理性高度近视;13:VKH)"
        )
    )
    quality_codes = parse_code_list(check_row.get("图像质量(1:合格;2:不合格)"))
    return lesion_codes, diagnosis_codes, quality_codes


def codes_to_text(ids: list[int], label_map: dict[int, str]) -> str:
    return "|".join(label_map[v] for v in ids if v in label_map)


def codes_to_str(ids: list[int]) -> str:
    return "|".join(str(v) for v in ids)


def aggregate_records_by_image_key(records: list[dict], dataset_version: str, start_index: int = 1) -> list[dict]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for record in records:
        grouped[record["image_key"]].append(record)

    aggregated: list[dict] = []
    sample_counter = start_index - 1
    for image_key in sorted(grouped):
        items = grouped[image_key]
        items = sorted(items, key=lambda x: (str(x.get("source_annotation_file", "")), str(x.get("patient_id", ""))))
        primary = items[0]

        lesion_codes = sorted({int(v) for item in items for v in parse_code_list(item.get("lesion_ids"))})
        diagnosis_codes = sorted({int(v) for item in items for v in parse_code_list(item.get("diagnosis_ids"))})
        quality_codes = sorted({int(v) for item in items for v in parse_code_list(item.get("quality_label"))})
        batch_ids = sorted({str(item.get("batch_id", "")) for item in items if item.get("batch_id")})
        source_annotation_files = sorted({str(item.get("source_annotation_file", "")) for item in items if item.get("source_annotation_file")})
        source_sheets = sorted({str(item.get("source_sheet", "")) for item in items if item.get("source_sheet")})
        patient_ids = sorted({str(item.get("patient_id", "")) for item in items if item.get("patient_id")})
        matched_paths = sorted({path for item in items for path in json.loads(item.get("all_matched_image_paths", "[]"))})
        lesion_signatures = {item.get("lesion_ids", "") for item in items}

        sample_counter += 1
        aggregated.append(
            {
                "sample_id": f"{dataset_version}_{sample_counter:06d}",
                "dataset_version": dataset_version,
                "batch_id": primary.get("batch_id", ""),
                "batch_ids": "|".join(batch_ids),
                "split": assign_split(image_key),
                "source_annotation_file": primary.get("source_annotation_file", ""),
                "source_annotation_files": json.dumps(source_annotation_files, ensure_ascii=False),
                "source_sheet": primary.get("source_sheet", ""),
                "source_sheets": "|".join(source_sheets),
                "patient_id": primary.get("patient_id", ""),
                "patient_ids": json.dumps(patient_ids, ensure_ascii=False),
                "study_id": primary.get("study_id", ""),
                "study_instance_uid": primary.get("study_instance_uid", ""),
                "series_instance_uid": primary.get("series_instance_uid", ""),
                "doctor_id": primary.get("doctor_id", ""),
                "image_key": image_key,
                "source_image_path": primary.get("source_image_path", ""),
                "all_matched_image_paths": json.dumps(matched_paths, ensure_ascii=False),
                "match_type": "unique" if len(matched_paths) == 1 else "multi",
                "quality_label": codes_to_str(quality_codes),
                "quality_text": "|".join(f"quality_{v}" for v in quality_codes),
                "diagnosis_ids": codes_to_str(diagnosis_codes),
                "diagnosis_texts": codes_to_text(diagnosis_codes, DIAGNOSIS_LABELS),
                "lesion_ids": codes_to_str(lesion_codes),
                "lesion_texts": codes_to_text(lesion_codes, LESION_LABELS),
                "num_lesions": len(lesion_codes),
                "annotation_record_count": len(items),
                "label_conflict": int(len(lesion_signatures) > 1),
                "linked_image_path": primary.get("linked_image_path", ""),
            }
        )
    return aggregated


def build_records_from_workbook(image_index: dict[str, list[str]], workbook_path: Path, batch_id: str, dataset_version: str) -> list[dict]:
    wb = load_workbook(workbook_path, read_only=False, data_only=True)
    lesion_sheet = choose_lesion_sheet(wb.sheetnames)
    if lesion_sheet is None:
        return []
    check_sheet = choose_check_sheet(wb.sheetnames, lesion_sheet)
    lesion_rows = worksheet_rows_by_patient(workbook_path, lesion_sheet)
    check_rows = worksheet_rows_by_patient(workbook_path, check_sheet) if check_sheet else {}

    records = []
    for patient_id, lesion_row in lesion_rows.items():
        image_key = normalize_key(patient_id)
        if image_key is None:
            continue
        matched_paths = sorted(image_index.get(image_key, []))
        if not matched_paths:
            continue
        check_row = check_rows.get(patient_id, {})
        lesion_codes, diagnosis_codes, quality_codes = extract_label_codes(lesion_row, check_row)
        records.append(
            {
                "sample_id": "",
                "dataset_version": dataset_version,
                "batch_id": batch_id,
                "split": assign_split(image_key),
                "source_annotation_file": str(workbook_path),
                "source_sheet": lesion_sheet,
                "patient_id": patient_id,
                "study_id": lesion_row.get("Study ID") or check_row.get("Study ID"),
                "study_instance_uid": lesion_row.get("Study Instance UID") or check_row.get("Study Instance UID"),
                "series_instance_uid": lesion_row.get("Series Instance UID") or check_row.get("Series Instance UID"),
                "doctor_id": lesion_row.get("Doctor ID") or check_row.get("Doctor ID"),
                "image_key": image_key,
                "source_image_path": matched_paths[0],
                "all_matched_image_paths": json.dumps(matched_paths, ensure_ascii=False),
                "match_type": "unique" if len(matched_paths) == 1 else "multi",
                "quality_label": codes_to_str(quality_codes),
                "quality_text": "|".join(f"quality_{v}" for v in quality_codes),
                "diagnosis_ids": codes_to_str(diagnosis_codes),
                "diagnosis_texts": codes_to_text(diagnosis_codes, DIAGNOSIS_LABELS),
                "lesion_ids": codes_to_str(lesion_codes),
                "lesion_texts": codes_to_text(lesion_codes, LESION_LABELS),
                "num_lesions": len(lesion_codes),
            }
        )
    return records


def build_dataset_records(
    image_root: Path,
    dataset_version: str,
    selected_annotations: dict[str, dict] | None = None,
) -> list[dict]:
    image_index = build_image_index(image_root)
    dataset_cfg = DATASET_SELECTIONS.get(dataset_version)
    if dataset_cfg is None and selected_annotations is None:
        raise KeyError(f"Unknown dataset version: {dataset_version}")
    selected = selected_annotations or dataset_cfg["selected_annotations"]
    records: list[dict] = []
    sample_counter = 0

    for batch_id, cfg in selected.items():
        workbook_path = Path(cfg["annotation_path"])
        if not workbook_path.exists():
            raise FileNotFoundError(f"Cannot find selected annotation file: {workbook_path}")
        lesion_rows = worksheet_rows_by_patient(workbook_path, cfg["lesion_sheet"])
        check_rows = worksheet_rows_by_patient(workbook_path, cfg["check_sheet"])

        for patient_id, lesion_row in lesion_rows.items():
            image_key = normalize_key(patient_id)
            if image_key is None:
                continue
            matched_paths = sorted(image_index.get(image_key, []))
            if not matched_paths:
                continue

            check_row = check_rows.get(patient_id, {})
            lesion_codes, diagnosis_codes, quality_codes = extract_label_codes(lesion_row, check_row)

            sample_counter += 1
            records.append(
                {
                    "sample_id": f"{dataset_version}_{sample_counter:06d}",
                    "dataset_version": dataset_version,
                    "batch_id": batch_id,
                    "split": assign_split(image_key),
                    "source_annotation_file": str(workbook_path),
                    "source_sheet": cfg["lesion_sheet"],
                    "patient_id": patient_id,
                    "study_id": lesion_row.get("Study ID") or check_row.get("Study ID"),
                    "study_instance_uid": lesion_row.get("Study Instance UID") or check_row.get("Study Instance UID"),
                    "series_instance_uid": lesion_row.get("Series Instance UID") or check_row.get("Series Instance UID"),
                    "doctor_id": lesion_row.get("Doctor ID") or check_row.get("Doctor ID"),
                    "image_key": image_key,
                    "source_image_path": matched_paths[0],
                    "all_matched_image_paths": json.dumps(matched_paths, ensure_ascii=False),
                    "match_type": "unique" if len(matched_paths) == 1 else "multi",
                    "quality_label": codes_to_str(quality_codes),
                    "quality_text": "|".join(f"quality_{v}" for v in quality_codes),
                    "diagnosis_ids": codes_to_str(diagnosis_codes),
                    "diagnosis_texts": codes_to_text(diagnosis_codes, DIAGNOSIS_LABELS),
                    "lesion_ids": codes_to_str(lesion_codes),
                    "lesion_texts": codes_to_text(lesion_codes, LESION_LABELS),
                    "num_lesions": len(lesion_codes),
                }
            )
    return records


def build_vision1_records(image_root: Path, annotation_root: Path | None = None) -> list[dict]:
    return build_dataset_records(
        image_root=image_root,
        dataset_version="vision1",
        selected_annotations=VISION1_SELECTED_ANNOTATIONS,
    )


def build_union_vision2_records(
    image_root: Path,
    vision1_manifest_path: Path,
    incoming_root: Path,
) -> list[dict]:
    image_index = build_image_index(image_root)

    with vision1_manifest_path.open("r", encoding="utf-8") as f:
        vision1_rows = list(csv.DictReader(f))
    base_records = aggregate_records_by_image_key(vision1_rows, dataset_version="vision2", start_index=1)
    base_keys = {record["image_key"] for record in base_records}

    incoming_records: list[dict] = []
    for dirpath, _, filenames in os.walk(incoming_root):
        for filename in sorted(filenames):
            if not filename.endswith(".xlsx") or filename.startswith("~$") or filename == ".xlsx":
                continue
            workbook_path = Path(dirpath) / filename
            batch_id = workbook_path.parent.name.split("_", 1)[0]
            workbook_records = build_records_from_workbook(
                image_index=image_index,
                workbook_path=workbook_path,
                batch_id=batch_id,
                dataset_version="vision2",
            )
            incoming_records.extend(record for record in workbook_records if record["image_key"] not in base_keys)

    new_records = aggregate_records_by_image_key(
        incoming_records,
        dataset_version="vision2",
        start_index=len(base_records) + 1,
    )
    return base_records + new_records


def build_union_vision3_records(
    vision2_manifest_path: Path,
    recovered_root: Path,
    missing_report_path: Path,
) -> list[dict]:
    with vision2_manifest_path.open("r", encoding="utf-8") as f:
        base_records = list(csv.DictReader(f))

    for record in base_records:
        record["dataset_version"] = "vision3"

    base_keys = {record["image_key"] for record in base_records}
    recovered_index = build_image_index(recovered_root)

    next_index = len(base_records) + 1
    new_records: list[dict] = []
    with missing_report_path.open("r", encoding="utf-8") as f:
        missing_rows = list(csv.DictReader(f))

    for row in missing_rows:
        image_key = str(row.get("image_key", "")).strip()
        if not image_key or image_key in base_keys:
            continue

        matched_paths = sorted(recovered_index.get(image_key, []))
        if not matched_paths:
            continue

        lesion_ids = str(row.get("lesion_ids", "")).strip()
        lesion_codes = parse_code_list(lesion_ids)
        diagnosis_ids = str(row.get("diagnosis_ids", "")).strip()
        diagnosis_codes = parse_code_list(diagnosis_ids)
        quality_label = str(row.get("quality_label", "")).strip()
        quality_codes = parse_code_list(quality_label)

        new_records.append(
            {
                "sample_id": f"vision3_{next_index:06d}",
                "dataset_version": "vision3",
                "batch_id": str(row.get("batch_id", "")).strip(),
                "split": assign_split(image_key),
                "source_annotation_file": str(row.get("source_annotation_file", "")).strip(),
                "source_sheet": str(row.get("source_sheet", "")).strip(),
                "patient_id": str(row.get("patient_id", "")).strip(),
                "study_id": "",
                "study_instance_uid": "",
                "series_instance_uid": "",
                "doctor_id": "",
                "image_key": image_key,
                "source_image_path": matched_paths[0],
                "all_matched_image_paths": json.dumps(matched_paths, ensure_ascii=False),
                "match_type": "unique" if len(matched_paths) == 1 else "multi",
                "quality_label": quality_label,
                "quality_text": str(row.get("quality_text", "")).strip(),
                "diagnosis_ids": diagnosis_ids,
                "diagnosis_texts": str(row.get("diagnosis_texts", "")).strip(),
                "lesion_ids": lesion_ids,
                "lesion_texts": str(row.get("lesion_texts", "")).strip(),
                "num_lesions": len(lesion_codes),
                "annotation_record_count": 1,
                "label_conflict": 0,
                "linked_image_path": "",
            }
        )
        next_index += 1

    return base_records + new_records


def write_manifest(records: list[dict], output_csv: Path) -> None:
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    if not records:
        raise ValueError("No records to write")
    with output_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(records[0].keys()))
        writer.writeheader()
        writer.writerows(records)


def write_label_maps(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "lesion_label_map.json").write_text(
        json.dumps(LESION_LABELS, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (output_dir / "diagnosis_label_map.json").write_text(
        json.dumps(DIAGNOSIS_LABELS, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
